#!/usr/bin/env python3
"""Generate the KSE Password Manager architecture deck as an editable PPTX."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE as Shape
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUT = Path(__file__).resolve().parents[1] / "docs" / "KSE_Password_Manager_Architecture.pptx"
W, H = Inches(13.333), Inches(7.5)

NAVY = "0B1F33"
BLUE = "0077B6"
SKY = "48CAE4"
TEAL = "00B4A6"
GREEN = "34B27B"
GOLD = "FFB703"
CORAL = "F05D5E"
INK = "102A43"
MUTED = "627D98"
LINE = "D9E2EC"
PALE = "F5F8FB"
WHITE = "FFFFFF"


def rgb(hex_value):
    hex_value = hex_value.replace("#", "")
    return RGBColor(*bytes.fromhex(hex_value))


def add_text(slide, text, x, y, w, h, size=18, color=INK, bold=False,
             align=PP_ALIGN.LEFT, font="Aptos", valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = rgb(color)
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    return box


def add_bullet_list(slide, items, x, y, w, h, size=17, color=INK, spacing=10):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = rgb(color)
        p.level = 0
        p.space_after = Pt(spacing)
        p.bullet = True
    return box


def add_rect(slide, x, y, w, h, fill=WHITE, line=LINE, radius=False):
    typ = Shape.ROUNDED_RECTANGLE if radius else Shape.RECTANGLE
    sh = slide.shapes.add_shape(typ, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(fill)
    sh.line.color.rgb = rgb(line)
    sh.line.width = Pt(0.8)
    return sh


def card(slide, title, body, x, y, w, h, accent=BLUE, body_size=14):
    add_rect(slide, x, y, w, h, WHITE, LINE, True)
    accent_shape = slide.shapes.add_shape(Shape.RECTANGLE, x, y, Inches(0.09), h)
    accent_shape.fill.solid(); accent_shape.fill.fore_color.rgb = rgb(accent); accent_shape.line.fill.background()
    add_text(slide, title, x + Inches(0.22), y + Inches(0.18), w - Inches(0.35), Inches(0.33), 16, INK, True)
    add_text(slide, body, x + Inches(0.22), y + Inches(0.60), w - Inches(0.40), h - Inches(0.74), body_size, MUTED)


def pill(slide, text, x, y, w, fill=PALE, color=BLUE):
    sh = add_rect(slide, x, y, w, Inches(0.30), fill, fill, True)
    add_text(slide, text, x, y + Inches(0.035), w, Inches(0.22), 10.5, color, True, PP_ALIGN.CENTER)
    return sh


def line(slide, x1, y1, x2, y2, color=LINE, width=1.6):
    connector = slide.shapes.add_connector(1, x1, y1, x2, y2)
    connector.line.color.rgb = rgb(color)
    connector.line.width = Pt(width)
    return connector


def base(prs, title, subtitle=None, number=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid(); bg.fore_color.rgb = rgb(WHITE)
    top = slide.shapes.add_shape(Shape.RECTANGLE, 0, 0, W, Inches(0.09))
    top.fill.solid(); top.fill.fore_color.rgb = rgb(BLUE); top.line.fill.background()
    add_text(slide, title, Inches(0.62), Inches(0.42), Inches(11.9), Inches(0.48), 26, NAVY, True)
    if subtitle:
        add_text(slide, subtitle, Inches(0.64), Inches(0.94), Inches(11.7), Inches(0.30), 12.5, MUTED)
    footer = "KSE Password Manager  •  Архітектурний огляд"
    add_text(slide, footer, Inches(0.62), Inches(7.13), Inches(7), Inches(0.18), 9.5, MUTED)
    if number is not None:
        add_text(slide, f"{number:02d}", Inches(12.16), Inches(7.10), Inches(0.55), Inches(0.20), 10, MUTED, True, PP_ALIGN.RIGHT)
    return slide


def title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill; bg.solid(); bg.fore_color.rgb = rgb(NAVY)
    bar = slide.shapes.add_shape(Shape.RECTANGLE, 0, 0, W, Inches(0.13))
    bar.fill.solid(); bar.fill.fore_color.rgb = rgb(SKY); bar.line.fill.background()
    add_text(slide, "KSE Password\nManager", Inches(0.8), Inches(1.2), Inches(7.0), Inches(1.65), 42, WHITE, True)
    add_text(slide, "Архітектура корпоративного менеджера паролів", Inches(0.84), Inches(3.05), Inches(6.8), Inches(0.40), 19, "B8D8EC")
    add_text(slide, "E2E-шифрування • Google Workspace SSO • командний доступ • аудит", Inches(0.84), Inches(3.65), Inches(7.9), Inches(0.30), 15, "B8D8EC")
    # abstract vault/key illustration
    add_rect(slide, Inches(9.05), Inches(1.38), Inches(2.75), Inches(3.7), "123958", "123958", True)
    add_rect(slide, Inches(9.67), Inches(2.22), Inches(1.5), Inches(1.4), NAVY, SKY, True)
    add_rect(slide, Inches(10.15), Inches(1.74), Inches(0.55), Inches(0.70), SKY, SKY, True)
    add_text(slide, "E2E", Inches(9.77), Inches(2.65), Inches(1.3), Inches(0.32), 20, WHITE, True, PP_ALIGN.CENTER)
    line(slide, Inches(9.10), Inches(5.46), Inches(11.85), Inches(5.46), SKY, 1)
    pill(slide, "React SPA", Inches(8.7), Inches(5.78), Inches(1.22), "123958", SKY)
    pill(slide, "Spring Boot", Inches(10.02), Inches(5.78), Inches(1.42), "123958", SKY)
    pill(slide, "Keycloak", Inches(11.53), Inches(5.78), Inches(1.10), "123958", SKY)
    add_text(slide, "Версія для огляду архітектури • 30 липня 2026", Inches(0.84), Inches(6.75), Inches(7.0), Inches(0.25), 11, "B8D8EC")


def slide_value(prs):
    s = base(prs, "Що вирішує застосунок", "Корпоративний менеджер доступів без передачі секретів у чатах", 2)
    card(s, "Єдине сховище", "Паролі, API credentials, SSH keys, картки, secure notes і TOTP в одному vault.", Inches(.72), Inches(1.60), Inches(3.75), Inches(2.0), BLUE, 16)
    card(s, "Контрольований доступ", "Прямий sharing, командні Collections, ролі й контроль відкликання доступу.", Inches(4.78), Inches(1.60), Inches(3.75), Inches(2.0), TEAL, 16)
    card(s, "Безпечна прозорість", "Audit log, аномалії, Google Workspace governance і policy-driven deployment.", Inches(8.84), Inches(1.60), Inches(3.75), Inches(2.0), GREEN, 16)
    add_text(s, "Ключовий принцип", Inches(.75), Inches(4.30), Inches(2.9), Inches(.35), 17, NAVY, True)
    add_text(s, "Backend зберігає зашифровані дані, загорнуті ключі, метадані та правила доступу — але не plaintext-паролі й не vault passphrase.", Inches(.75), Inches(4.78), Inches(10.9), Inches(.74), 24, INK, True)
    pill(s, "API-first", Inches(.75), Inches(5.85), Inches(1.18), "E8F5FD", BLUE)
    pill(s, "Zero-knowledge vault", Inches(2.10), Inches(5.85), Inches(1.95), "E5F7F3", TEAL)
    pill(s, "Google Workspace", Inches(4.23), Inches(5.85), Inches(1.84), "EDF8F1", GREEN)
    pill(s, "Chrome extension", Inches(6.25), Inches(5.85), Inches(1.80), "FFF5DB", "A66A00")


def slide_system(prs):
    s = base(prs, "Карта системи", "Клієнт шифрує; API керує політикою; зовнішні сервіси дають identity та encrypted storage", 3)
    y = Inches(2.15)
    boxes = [
        ("React SPA\nWebCrypto", "browser", BLUE),
        ("Chrome\nExtension", "browser", TEAL),
        ("nginx\nTLS edge", "public ingress", GOLD),
        ("Spring Boot\nREST API", "policy + audit", BLUE),
        ("PostgreSQL\nFlyway", "ciphertext + policy", "5C6BC0"),
        ("Keycloak\nGoogle SSO", "OIDC / JWT", GREEN),
        ("Google Workspace\nDirectory + Drive", "governance + files", "2B8A3E"),
    ]
    xs = [0.65, 2.20, 4.00, 5.78, 7.65, 9.52, 11.20]
    widths = [1.28, 1.28, 1.42, 1.52, 1.48, 1.38, 1.45]
    for i, ((title, sub, accent), x, w) in enumerate(zip(boxes, xs, widths)):
        add_rect(s, Inches(x), y, Inches(w), Inches(1.55), PALE, LINE, True)
        add_rect(s, Inches(x), y, Inches(w), Inches(.11), accent, accent, True)
        add_text(s, title, Inches(x+.10), y+Inches(.29), Inches(w-.2), Inches(.55), 14, INK, True, PP_ALIGN.CENTER)
        add_text(s, sub, Inches(x+.08), y+Inches(1.02), Inches(w-.16), Inches(.22), 9.8, MUTED, False, PP_ALIGN.CENTER)
        if i < len(boxes)-1:
            line(s, Inches(x+w), y+Inches(.78), Inches(xs[i+1]), y+Inches(.78), SKY, 1.8)
    # lower interactions
    line(s, Inches(6.55), Inches(4.28), Inches(6.55), Inches(5.05), TEAL, 1.6)
    line(s, Inches(6.55), Inches(5.05), Inches(8.4), Inches(5.05), TEAL, 1.6)
    add_text(s, "HTTPS", Inches(3.27), Inches(1.88), Inches(.65), Inches(.20), 10, MUTED, True, PP_ALIGN.CENTER)
    add_text(s, "JWT", Inches(9.05), Inches(1.88), Inches(.65), Inches(.20), 10, MUTED, True, PP_ALIGN.CENTER)
    card(s, "Серверна межа довіри", "REST API перевіряє identity, governance, ролі та audit. Він не виконує розшифрування vault-даних.", Inches(.82), Inches(4.72), Inches(4.55), Inches(1.42), BLUE, 14)
    card(s, "Дані", "PostgreSQL: cipher­text, envelopes, metadata, permissions, audit. Drive: лише encrypted attachments.", Inches(5.67), Inches(4.72), Inches(3.18), Inches(1.42), "5C6BC0", 13)
    card(s, "Identity і directory", "Keycloak брокерує Google SSO. Workspace snapshot — реальна межа corporate access.", Inches(9.15), Inches(4.72), Inches(3.18), Inches(1.42), GREEN, 13)


def slide_crypto(prs):
    s = base(prs, "E2E-шифрування: ключі не покидають клієнт", "Per-item key wrapping робить sharing багатокористувацьким без довіри до backend", 4)
    steps = [
        ("1", "Vault passphrase", "У браузері виводить KEK через Argon2id або legacy PBKDF2."),
        ("2", "RSA keypair", "Private key AES-GCM encrypted KEK; public key доступний для sharing."),
        ("3", "Item key", "Кожен запис має власний AES-256 key; поля шифруються AES-GCM."),
        ("4", "Wrapped envelopes", "Item key RSA-OAEP загорнутий для кожного отримувача доступу."),
    ]
    for i, (n, title, body) in enumerate(steps):
        x = Inches(.75 + i*3.12)
        add_rect(s, x, Inches(1.75), Inches(2.73), Inches(2.16), WHITE, LINE, True)
        circ = s.shapes.add_shape(Shape.OVAL, x+Inches(.18), Inches(1.95), Inches(.45), Inches(.45))
        circ.fill.solid(); circ.fill.fore_color.rgb = rgb([BLUE, TEAL, "5C6BC0", GREEN][i]); circ.line.fill.background()
        add_text(s, n, x+Inches(.18), Inches(2.045), Inches(.45), Inches(.16), 11, WHITE, True, PP_ALIGN.CENTER)
        add_text(s, title, x+Inches(.18), Inches(2.60), Inches(2.30), Inches(.36), 16, INK, True)
        add_text(s, body, x+Inches(.18), Inches(3.08), Inches(2.30), Inches(.58), 12, MUTED)
        if i < 3:
            line(s, x+Inches(2.73), Inches(2.83), x+Inches(3.03), Inches(2.83), SKY, 2.0)
    add_rect(s, Inches(.75), Inches(4.55), Inches(11.85), Inches(1.25), "E8F5FD", "E8F5FD", True)
    add_text(s, "Collections", Inches(1.00), Inches(4.80), Inches(1.45), Inches(.28), 18, BLUE, True)
    add_text(s, "Collection key RSA-wrapped для кожного члена → AES-wrap item keys у колекції. Додати учасника = створити один новий envelope, а не перешифровувати всі записи.", Inches(2.50), Inches(4.76), Inches(9.3), Inches(.50), 15, INK)
    add_text(s, "Revoke: клієнт спочатку ротує item або collection key, потім API прибирає доступ. Це захищає майбутні читання, але не стирає дані, уже скопійовані отримувачем.", Inches(.85), Inches(6.22), Inches(11.6), Inches(.40), 13.5, MUTED)


def slide_auth(prs):
    s = base(prs, "Вхід і контроль доступу", "Keycloak встановлює особу; локальний Workspace snapshot ухвалює рішення про доступ", 5)
    stages = [
        ("Google SSO", "Користувач автентифікується корпоративним Google account."),
        ("Keycloak", "Брокерує IdP, видає OIDC access token; passwordless passkey доступний як альтернатива."),
        ("JWT validation", "Spring Resource Server перевіряє signature та issuer."),
        ("Governance", "Filter звіряє email з local Workspace snapshot і встановлює ефективну роль."),
    ]
    for i, (title, body) in enumerate(stages):
        x = Inches(.72 + i*3.12)
        col = [GREEN, "5C6BC0", BLUE, TEAL][i]
        add_rect(s, x, Inches(1.80), Inches(2.75), Inches(1.90), WHITE, LINE, True)
        pill(s, f"{i+1}", x+Inches(.18), Inches(2.02), Inches(.34), col, WHITE)
        add_text(s, title, x+Inches(.18), Inches(2.53), Inches(2.35), Inches(.30), 16, INK, True)
        add_text(s, body, x+Inches(.18), Inches(2.98), Inches(2.33), Inches(.48), 12.2, MUTED)
        if i < 3: line(s, x+Inches(2.75), Inches(2.75), x+Inches(3.05), Inches(2.75), SKY, 2)
    card(s, "Corporate roles", "USER • ADMIN • SECURITY\nДжерело: Workspace snapshot, не JWT realm role.", Inches(.85), Inches(4.45), Inches(3.55), Inches(1.33), BLUE, 13)
    card(s, "Record & collection roles", "OWNER / EDITOR / VIEWER / PASSKEY_ONLY\nMANAGER / EDITOR / VIEWER", Inches(4.72), Inches(4.45), Inches(3.55), Inches(1.33), TEAL, 13)
    card(s, "Важлива межа", "Адмін читає audit, але не отримує auto-access до чужих ciphertext-ів або ключів.", Inches(8.58), Inches(4.45), Inches(3.55), Inches(1.33), GREEN, 13)
    add_text(s, "Degraded mode можливий лише на порожньому snapshot: in-domain SSO користувач отримує тільки USER. Його можна вимкнути для fail-closed deployment.", Inches(.85), Inches(6.30), Inches(11.4), Inches(.32), 12.5, MUTED)


def slide_backend(prs):
    s = base(prs, "Backend і дані", "Spring Boot 3 / Java 21: доменні модулі, REST API, Flyway та PostgreSQL", 6)
    modules = [
        ("auth", "JWT, governance, rate limit, request context", BLUE),
        ("vault + keys", "CRUD, versions, import, encrypted keys, recovery", "5C6BC0"),
        ("sharing + collections", "Permissions, members, roles, key rotation", TEAL),
        ("directory + google", "Workspace snapshot, DWD token, sync", GREEN),
        ("attachments + external", "Drive stream, public encrypted package", GOLD),
        ("passkeys + audit", "Extension-only credentials, events, anomalies", CORAL),
    ]
    for idx, (title, body, col) in enumerate(modules):
        row, c = divmod(idx, 3)
        card(s, title, body, Inches(.75+c*4.12), Inches(1.55+row*1.55), Inches(3.70), Inches(1.23), col, 12.5)
    add_rect(s, Inches(.75), Inches(5.02), Inches(11.88), Inches(1.05), PALE, LINE, True)
    add_text(s, "PostgreSQL + Flyway", Inches(1.00), Inches(5.24), Inches(2.35), Inches(.25), 17, NAVY, True)
    add_text(s, "V1–V23 migrations • Hibernate ddl-auto=validate • ciphertext, key envelopes, access policy, Workspace snapshot, audit і metadata", Inches(3.33), Inches(5.22), Inches(8.65), Inches(.36), 14, MUTED)
    add_text(s, "REST API-first: вебклієнт і extension — клієнти одного контракту; ця межа залишає шлях до mobile client без дублювання policy.", Inches(.85), Inches(6.40), Inches(11.6), Inches(.25), 12.5, MUTED)


def slide_features(prs):
    s = base(prs, "Функціональні сценарії", "Продуктові можливості розміщуються поверх однієї криптографічної та policy-моделі", 7)
    cards = [
        ("Vault", "Typed records: LOGIN, API credentials, SSH key, card, secure note; versions, tags, CSV import/export.", BLUE),
        ("Sharing", "User/group permissions; Collections для реального командного E2E-доступу; client-side key rotation.", TEAL),
        ("External share", "Окремий encrypted snapshot, TTL або one-time read; key залишається у URL fragment.", GOLD),
        ("Attachments", "Encrypted file stream у per-owner Drive appDataFolder; metadata та checksum у Postgres.", GREEN),
        ("Security dashboard", "Після unlock: weak/reused/old/breached. HIBP бачить тільки SHA-1 prefix.", CORAL),
        ("Audit", "Create/share/revoke/copy та security anomalies для ADMIN/SECURITY.", "5C6BC0"),
    ]
    for i, (title, body, col) in enumerate(cards):
        r, c = divmod(i, 3)
        card(s, title, body, Inches(.76+c*4.10), Inches(1.58+r*1.78), Inches(3.70), Inches(1.48), col, 13)
    add_text(s, "Local-only records лишаються тільки у browser localStorage під KEK: вони не доступні на інших пристроях, не можуть бути shared і не потрапляють у server backup.", Inches(.85), Inches(5.52), Inches(11.3), Inches(.46), 14, INK, True)
    pill(s, "Client-side crypto", Inches(.85), Inches(6.25), Inches(1.65), "E8F5FD", BLUE)
    pill(s, "Policy-enforced API", Inches(2.68), Inches(6.25), Inches(1.75), "E5F7F3", TEAL)
    pill(s, "Auditability", Inches(4.62), Inches(6.25), Inches(1.22), "FFF5DB", "A66A00")


def slide_extension(prs):
    s = base(prs, "Chrome extension: окремий клієнт із вузькою довірою", "Manifest V3, origin-aware autofill і extension-only passkey material", 8)
    add_rect(s, Inches(.78), Inches(1.64), Inches(2.35), Inches(3.80), "E8F5FD", "D5EAF7", True)
    add_text(s, "Content script", Inches(1.02), Inches(1.96), Inches(1.85), Inches(.30), 18, BLUE, True, PP_ALIGN.CENTER)
    add_text(s, "Показує picker, заповнює лише вибрані поля, сканує QR TOTP. Не отримує весь vault.", Inches(1.02), Inches(2.53), Inches(1.85), Inches(1.05), 13, INK, False, PP_ALIGN.CENTER)
    add_text(s, "Strict origin match", Inches(1.00), Inches(4.56), Inches(1.9), Inches(.25), 11, MUTED, True, PP_ALIGN.CENTER)
    line(s, Inches(3.15), Inches(3.48), Inches(4.03), Inches(3.48), SKY, 2.2)
    add_rect(s, Inches(4.05), Inches(1.64), Inches(2.35), Inches(3.80), "E5F7F3", "CDEDE6", True)
    add_text(s, "Background worker", Inches(4.28), Inches(1.96), Inches(1.90), Inches(.30), 18, TEAL, True, PP_ALIGN.CENTER)
    add_text(s, "Пам'яттєва unlock-сесія, OIDC PKCE, API calls, TOTP, clipboard timeout і WebAuthn proxy.", Inches(4.30), Inches(2.53), Inches(1.86), Inches(1.05), 13, INK, False, PP_ALIGN.CENTER)
    add_text(s, "Lock → clean memory", Inches(4.28), Inches(4.56), Inches(1.92), Inches(.25), 11, MUTED, True, PP_ALIGN.CENTER)
    line(s, Inches(6.42), Inches(3.48), Inches(7.30), Inches(3.48), SKY, 2.2)
    add_rect(s, Inches(7.32), Inches(1.64), Inches(2.35), Inches(3.80), "FFF5DB", "FCE6A7", True)
    add_text(s, "Web app bridge", Inches(7.57), Inches(1.96), Inches(1.85), Inches(.30), 18, "A66A00", True, PP_ALIGN.CENTER)
    add_text(s, "App виконує vault WebAuthn ceremony; origin перевіряється на обох сторонах bridge.", Inches(7.56), Inches(2.53), Inches(1.87), Inches(1.05), 13, INK, False, PP_ALIGN.CENTER)
    add_text(s, "Unlock: app → extension", Inches(7.47), Inches(4.56), Inches(2.02), Inches(.25), 11, MUTED, True, PP_ALIGN.CENTER)
    line(s, Inches(9.69), Inches(3.48), Inches(10.57), Inches(3.48), SKY, 2.2)
    add_rect(s, Inches(10.59), Inches(1.64), Inches(1.96), Inches(3.80), "EDF8F1", "D4F0DF", True)
    add_text(s, "API", Inches(10.83), Inches(1.96), Inches(1.48), Inches(.30), 18, GREEN, True, PP_ALIGN.CENTER)
    add_text(s, "Passkey endpoints додатково перевіряють OAuth client ID та audience.", Inches(10.83), Inches(2.53), Inches(1.45), Inches(1.05), 13, INK, False, PP_ALIGN.CENTER)
    add_text(s, "Chrome 115+", Inches(10.83), Inches(4.56), Inches(1.45), Inches(.25), 11, MUTED, True, PP_ALIGN.CENTER)
    add_text(s, "Сайт ніколи не одержує full vault: extension передає контенту тільки дані для конкретної user-selected дії.", Inches(.92), Inches(6.15), Inches(11.4), Inches(.34), 13.5, MUTED)


def slide_google(prs):
    s = base(prs, "Google Workspace: governance без залежності від live lookup", "Локальний snapshot зменшує blast radius, навантаження й наслідки outage", 9)
    # credential modes
    card(s, "A. WIF (рекомендовано поза GCP)", "Keycloak client_credentials → Google STS → IAM signJwt. Без downloaded key.", Inches(.75), Inches(1.55), Inches(3.77), Inches(1.40), BLUE, 12.5)
    card(s, "B. Attached service account", "GCE/GKE metadata ADC → IAM signJwt. Keyless та self-renewing.", Inches(4.78), Inches(1.55), Inches(3.77), Inches(1.40), TEAL, 12.5)
    card(s, "C. JSON key — тільки fallback", "Mount read-only, не класти в image/git, регулярно ротути.", Inches(8.81), Inches(1.55), Inches(3.77), Inches(1.40), GOLD, 12.5)
    add_text(s, "Sync і outage behavior", Inches(.80), Inches(3.42), Inches(2.55), Inches(.30), 18, NAVY, True)
    flow = [("Directory API", GREEN), ("15 хв sync", BLUE), ("Local snapshot", TEAL), ("Governance + picker", "5C6BC0")]
    for i, (txt, col) in enumerate(flow):
        x = Inches(.85+i*3.05)
        add_rect(s, x, Inches(4.02), Inches(2.50), Inches(.82), WHITE, LINE, True)
        add_text(s, txt, x, Inches(4.27), Inches(2.50), Inches(.25), 14, col, True, PP_ALIGN.CENTER)
        if i < 3: line(s, x+Inches(2.50), Inches(4.43), x+Inches(2.97), Inches(4.43), SKY, 2)
    add_bullet_list(s, [
        "Recipient search зазвичай працює зі snapshot, а не з Google на кожний символ.",
        "Перший sync retry: 1 хв; після успіху — 15 хв. Старий snapshot зберігається при помилці.",
        "Provisional access — тільки USER, тільки для in-domain account і тільки доки snapshot порожній.",
    ], Inches(.90), Inches(5.32), Inches(11.3), Inches(1.10), 13, MUTED, 4)


def slide_deploy(prs):
    s = base(prs, "Production deployment", "Єдиний публічний edge, окремі origins для app і Keycloak, внутрішні сервіси без відкритих портів", 10)
    add_rect(s, Inches(.78), Inches(1.55), Inches(2.42), Inches(4.25), "E8F5FD", "D5EAF7", True)
    add_text(s, "Internet", Inches(1.1), Inches(1.90), Inches(1.80), Inches(.30), 19, BLUE, True, PP_ALIGN.CENTER)
    add_text(s, "https://pm.example.com\nhttps://auth.pm.example.com", Inches(1.03), Inches(2.62), Inches(1.95), Inches(.60), 13, INK, False, PP_ALIGN.CENTER)
    line(s, Inches(3.22), Inches(3.70), Inches(4.12), Inches(3.70), SKY, 2.3)
    add_rect(s, Inches(4.14), Inches(1.55), Inches(2.23), Inches(4.25), "FFF5DB", "FCE6A7", True)
    add_text(s, "nginx", Inches(4.48), Inches(1.92), Inches(1.55), Inches(.30), 20, "A66A00", True, PP_ALIGN.CENTER)
    add_bullet_list(s, ["TLS termination", "HSTS / CSP / headers", "Health only public", "Attachment streaming"], Inches(4.53), Inches(2.55), Inches(1.48), Inches(1.70), 12, INK, 4)
    line(s, Inches(6.39), Inches(3.70), Inches(7.30), Inches(3.70), SKY, 2.3)
    # internal grouped
    add_rect(s, Inches(7.32), Inches(1.35), Inches(5.25), Inches(4.75), PALE, LINE, True)
    add_text(s, "Внутрішня Docker-мережа", Inches(7.60), Inches(1.63), Inches(4.65), Inches(.27), 14, MUTED, True, PP_ALIGN.CENTER)
    for title, sub, x, y, col in [
        ("Spring Boot", "API + SPA", 7.75, 2.28, BLUE),
        ("PostgreSQL", "data volume", 10.12, 2.28, "5C6BC0"),
        ("Keycloak", "OIDC / realm", 7.75, 4.12, GREEN),
        ("keycloak-init", "one-shot setup", 10.12, 4.12, TEAL),
    ]:
        add_rect(s, Inches(x), Inches(y), Inches(1.98), Inches(1.22), WHITE, LINE, True)
        add_text(s, title, Inches(x+.1), Inches(y+.30), Inches(1.78), Inches(.24), 14, col, True, PP_ALIGN.CENTER)
        add_text(s, sub, Inches(x+.1), Inches(y+.70), Inches(1.78), Inches(.20), 10.5, MUTED, False, PP_ALIGN.CENTER)
    add_text(s, "Secrets: Secret Manager або protected environment file. Ніколи не image/git. Перевірки: health, login/logout, decrypt, passkey, external share.", Inches(.85), Inches(6.32), Inches(11.4), Inches(.34), 13.5, MUTED)


def slide_security_ops(prs):
    s = base(prs, "Security і операційна готовність", "Defense-in-depth поєднує edge, API, encryption, audit і runbooks", 11)
    left = [
        ("Edge", "TLS, HSTS, CSP, nosniff, X-Frame-Options; Swagger/Actuator blocked externally."),
        ("API", "Stateless bearer tokens, rate limits, trusted proxy IP model, internal audit endpoint."),
        ("Crypto", "Passphrase та raw keys не зберігаються сервером; client-side AES-GCM / RSA-OAEP."),
    ]
    right = [
        ("Backups", "Щоденний PostgreSQL dump, off-host storage, Keycloak realm/DB backup, monthly restore test."),
        ("Observability", "ECS JSON logs із actorEmail/requestId; Prometheus внутрішньо; health endpoint."),
        ("Alerts", "up=0, 5xx, 429, Hikari pending; audit anomalies для copy/share/revoke."),
    ]
    for i, (t, b) in enumerate(left): card(s, t, b, Inches(.78), Inches(1.54+i*1.45), Inches(5.65), Inches(1.15), [BLUE, TEAL, "5C6BC0"][i], 13)
    for i, (t, b) in enumerate(right): card(s, t, b, Inches(6.90), Inches(1.54+i*1.45), Inches(5.65), Inches(1.15), [GOLD, GREEN, CORAL][i], 13)
    add_rect(s, Inches(.78), Inches(6.05), Inches(11.78), Inches(.62), "FFF5DB", "FCE6A7", True)
    add_text(s, "Відновлення: server backup повертає server-synced cipher­text і envelopes, але не local-only records і не забутий passphrase.", Inches(1.00), Inches(6.23), Inches(11.3), Inches(.22), 13, "7A4B00", True)


def slide_limits(prs):
    s = base(prs, "Обмеження, ризики та наступні кроки", "Архітектура фіксує не лише можливості, а й чесні межі поточної реалізації", 12)
    card(s, "MFA", "Passkey підтримується, але не обов'язково перевіряється при кожному повторному Google login.", Inches(.78), Inches(1.58), Inches(3.72), Inches(1.43), CORAL, 13)
    card(s, "Group sharing", "Прямий group-share не створює envelopes автоматично; для team E2E потрібні Collections.", Inches(4.80), Inches(1.58), Inches(3.72), Inches(1.43), GOLD, 13)
    card(s, "Audit delivery", "HTTP endpoint готовий, але Keycloak Event Listener SPI ще треба розгорнути і перевірити.", Inches(8.82), Inches(1.58), Inches(3.72), Inches(1.43), "5C6BC0", 13)
    card(s, "Scale", "Rate limiter in-memory/per-node; горизонтальне масштабування потребує shared store (Redis/Bucket4j).", Inches(.78), Inches(3.45), Inches(3.72), Inches(1.43), BLUE, 13)
    card(s, "Recovery", "Server-side key escrow відсутній: offboarding вимагає попереднього sharing/Collections.", Inches(4.80), Inches(3.45), Inches(3.72), Inches(1.43), TEAL, 13)
    card(s, "Roadmap", "Mobile client, Firefox/Safari, enterprise policies та richer security monitoring.", Inches(8.82), Inches(3.45), Inches(3.72), Inches(1.43), GREEN, 13)
    add_rect(s, Inches(.78), Inches(5.65), Inches(11.75), Inches(.65), NAVY, NAVY, True)
    add_text(s, "Головний висновок: застосунок тримає централізовану політику та аудит, не перетворюючи backend на власника секретів.", Inches(1.00), Inches(5.84), Inches(11.3), Inches(.26), 15, WHITE, True, PP_ALIGN.CENTER)


def main():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    title_slide(prs)
    slide_value(prs)
    slide_system(prs)
    slide_crypto(prs)
    slide_auth(prs)
    slide_backend(prs)
    slide_features(prs)
    slide_extension(prs)
    slide_google(prs)
    slide_deploy(prs)
    slide_security_ops(prs)
    slide_limits(prs)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
